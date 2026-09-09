"""Word / Excel 的共享只读解析：聊天展示与 RAG 建库共用，不含用户话术。"""
from pathlib import Path


def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def _is_heading(style_name: str) -> bool:
    """兼容 Word 中文与英文内置标题样式。"""
    style = (style_name or "").strip().casefold()
    return style.startswith("heading") or style.startswith("标题")


def extract_docx_sections(path: Path) -> list[dict]:
    """保留 Word 标题和表格，返回 [{title, text}]。"""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    sections: list[dict] = []
    title, lines = "", []

    def flush() -> None:
        nonlocal lines
        text = "\n".join(lines).strip()
        lines = []
        if text:
            sections.append({"title": title, "text": text})

    # 按 body 原顺序读取段落和表格，避免把表格统一追加到全文末尾。
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            para = Paragraph(child, doc)
            text = (para.text or "").strip()
            if not text:
                continue
            if _is_heading(getattr(para.style, "name", "")):
                flush()
                title = text
            else:
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells = [_cell_text(cell.text) for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                lines.append("表格：\n" + "\n".join(rows))
    flush()
    return sections


def extract_xlsx_sections(path: Path, rows_per_chunk: int = 100) -> list[dict]:
    """按工作表、表头和行批次切 Excel；每一批都重复表头。"""
    from openpyxl import load_workbook

    if rows_per_chunk < 1:
        raise ValueError("rows_per_chunk 必须大于 0")
    wb = load_workbook(str(path), data_only=True, read_only=True)
    sections: list[dict] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(value) for value in row]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        header, data_rows = rows[0], rows[1:]
        header_text = " | ".join(header)
        # 只有表头时仍让知识库能回答“该表有哪些字段”。
        groups = [data_rows[i:i + rows_per_chunk] for i in range(0, len(data_rows), rows_per_chunk)] or [[]]
        for number, group in enumerate(groups, start=1):
            lines = [f"表头：{header_text}"]
            lines.extend(" | ".join(row) for row in group)
            suffix = f"（第 {number} 批）" if len(groups) > 1 else ""
            sections.append({"title": f"表：{ws.title}{suffix}", "text": "\n".join(lines)})
    return sections


def docx_display_text(path: Path) -> tuple[int, str]:
    """给聊天工具的 Word 展示文本；返回段落/表格段数和正文。"""
    sections = extract_docx_sections(path)
    text = "\n".join(
        "\n".join(part for part in (section["title"], section["text"]) if part)
        for section in sections
    )
    return len(sections), text


def xlsx_display_text(path: Path, max_rows_per_sheet: int = 200) -> str:
    """给聊天工具的 Excel 展示文本；保留原有每表 200 行上限。"""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"【表：{ws.title}】")
        shown = 0
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(value) for value in row]
            if any(cells):
                out.append(" | ".join(cells))
                shown += 1
            if shown >= max_rows_per_sheet:
                out.append(f"…（行数较多，仅显示前 {max_rows_per_sheet} 行）")
                break
    return "\n".join(out)
