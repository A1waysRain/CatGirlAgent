# -*- coding: utf-8 -*-
"""隔离 APPDATA 验证 M2：append_file 给已有 pptx 追加页（主题继承）。"""
import os, sys, tempfile, shutil
tmp = tempfile.mkdtemp(prefix="catgirl_appd_")
os.environ["APPDATA"] = tmp
sys.path.insert(0, ".")
import backend.tools as T
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

def check(label, cond, extra=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {label}" + (f"  {extra}" if extra else ""))
    return cond

ok = True
pptx_p = os.path.join(tmp, "水务汇报.pptx")
decks = [
    {"layout": "title", "title": "水务运营汇报", "subtitle": "2026"},
    {"layout": "content", "title": "现状", "bullets": ["水源调度", "- 一级要点"]},
]
r = T.tool_create_pptx(pptx_p, decks, confirmed=True)
ok &= check("先建2页水务deck(主题ocean)", "已经生成" in r, r[:40])

# 追加一页（第一行当标题，其余当要点）
r = T.tool_append_file(pptx_p, "新增章节\n要点甲\n- 要点乙", confirmed=True)
ok &= check("追加成功", "末尾加了一页" in r, r[:40])

prs = Presentation(pptx_p)
ok &= check("页数 2→3", len(prs.slides) == 3)

# 新页 = 最后一页：标题=第一行，含要点
last = prs.slides[-1]
texts = []
for sh in last.shapes:
    if sh.has_text_frame:
        texts.append(sh.text_frame.text)
joined = "|".join(texts)
ok &= check("新页标题=新增章节", "新增章节" in joined)
ok &= check("新页含要点甲", "要点甲" in joined)
ok &= check("新页含二级要点", "–" in joined and "要点乙" in joined)

# 新页主题继承 ocean（标题栏 = darken(0F9D9D,0.78)，圆点 = 0F9D9D）
hdr_fills, dot_colors = [], []
for sh in last.shapes:
    try:
        if sh.auto_shape_type == MSO_SHAPE.RECTANGLE and sh.fill.type is not None:
            hdr_fills.append(str(sh.fill.fore_color.rgb))
    except ValueError:
        pass
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for run in p.runs:
                if run.text.startswith("●"):
                    dot_colors.append(str(run.font.color.rgb))
ok &= check("新页标题栏是ocean深色", any(c != "FFFFFF" for c in hdr_fills), f"fills={hdr_fills[:2]}")
ok &= check("新页圆点是ocean主题色", dot_colors and dot_colors[0] == "0F9D9D", f"dots={dot_colors[:1]}")
# 新页脚 3/3
ok &= check("新页脚 3/3", any("3 / 3" in t for t in texts))

# 旧页不动
first = prs.slides[0]
old_txt = [sh.text_frame.text for sh in first.shapes if sh.has_text_frame]
ok &= check("旧页仍在", any("水务运营汇报" in t for t in old_txt))

# 空内容拒绝
r = T.tool_append_file(pptx_p, "   \n ", confirmed=True)
ok &= check("空内容拒绝", "空" in r)

# 坏文件拒绝
bad = os.path.join(tmp, "bad.pptx")
open(bad, "w", encoding="utf-8").write("不是pptx")
r = T.tool_append_file(bad, "内容", confirmed=True)
ok &= check("坏文件拒绝", "打不开" in r)

shutil.rmtree(tmp, ignore_errors=True)
print("ALL PASS" if ok else "HAS FAILURES")
