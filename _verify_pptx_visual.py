# -*- coding: utf-8 -*-
"""验证 create_pptx 的视觉结构：背景填充/标题栏/彩色圆点/装饰形状。"""
import os, sys, tempfile, shutil
tmp = tempfile.mkdtemp(prefix="catgirl_pptvis_")
os.environ["APPDATA"] = tmp
sys.path.insert(0, ".")
import backend.tools as T
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

def check(label, cond, extra=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {label}" + (f"  {extra}" if extra else ""))
    return cond

pptx_p = os.path.join(tmp, "v.pptx")
slides = [
    {"layout": "title", "title": "封面标题", "subtitle": "副题"},
    {"layout": "section", "title": "第一章"},
    {"layout": "content", "title": "要点页", "bullets": ["一级要点", "- 二级要点"]},
    {"layout": "end", "title": "谢谢"},
]
r = T.tool_create_pptx(pptx_p, slides, confirmed=True)
prs = Presentation(pptx_p)
slides_iter = list(prs.slides)
ok = True

# ① 封面：背景有填充、非默认白
s0 = slides_iter[0]
bg_fill = s0.background.fill
has_bg = bg_fill.type is not None and str(bg_fill.fore_color.rgb) != "FFFFFF"
ok &= check("封面背景=主题色填充", has_bg, str(bg_fill.fore_color.rgb))
title_texts = [sh.text_frame.text for sh in s0.shapes if sh.has_text_frame]
ok &= check("封面标题文字", any("封面标题" in t for t in title_texts))

# ② 章节页：大数字 01 存在
s1 = slides_iter[1]
sec_texts = [sh.text_frame.text for sh in s1.shapes if sh.has_text_frame]
ok &= check("章节大数字01", any("01" in t for t in sec_texts))
ok &= check("章节标题", any("第一章" in t for t in sec_texts))

# ③ 内容页：有标题栏矩形 + 背景非白 + 彩色圆点 run
s2 = slides_iter[2]
rect_fills = []
bullet_dot_colors = []
for sh in s2.shapes:
    try:
        ast = sh.auto_shape_type
    except ValueError:
        ast = None
    if ast in (MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE):
        if sh.fill.type is not None:
            rect_fills.append(str(sh.fill.fore_color.rgb))
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for run in p.runs:
                if run.text.startswith("●"):
                    bullet_dot_colors.append(str(run.font.color.rgb))
ok &= check("内容页有标题栏(深粉矩形)", len(rect_fills) > 0, f"fills={rect_fills[:2]}")
ok &= check("标题栏是深一档主题色", any(c != "FFFFFF" for c in rect_fills))
ok &= check("要点有彩色圆点●", len(bullet_dot_colors) >= 1, f"colors={bullet_dot_colors[:2]}")
ok &= check("圆点是主题色", bullet_dot_colors and bullet_dot_colors[0] != "2B2B2B")

# ④ 内容页背景非纯白（浅粉底）
ok &= check("内容页背景=浅粉", str(s2.background.fill.fore_color.rgb) != "FFFFFF")

# ⑤ 页脚页码
s2_texts = [sh.text_frame.text for sh in s2.shapes if sh.has_text_frame]
ok &= check("内容页有页脚 3/4", any("3 / 4" in t for t in s2_texts))
ok &= check("页脚品牌", any("猫娘来咯" in t for t in s2_texts))

# ⑥ 结尾页背景=主题色
ok &= check("结尾背景=主题色", str(slides_iter[3].background.fill.fore_color.rgb) != "FFFFFF")

shutil.rmtree(tmp, ignore_errors=True)
print("ALL PASS" if ok else "HAS FAILURES")
