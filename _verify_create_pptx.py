# -*- coding: utf-8 -*-
"""隔离 APPDATA 验证 M1 create_pptx。"""
import os, sys, tempfile, shutil
tmp = tempfile.mkdtemp(prefix="catgirl_pptx_")
os.environ["APPDATA"] = tmp
sys.path.insert(0, ".")
import backend.tools as T
from pptx import Presentation
from pptx.oxml.ns import qn
from pathlib import Path

def check(label, cond, extra=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {label}" + (f"  {extra}" if extra else ""))
    return cond

ok = True
pptx_p = os.path.join(tmp, "方案汇报.pptx")
slides = [
    {"layout": "title", "title": "水务平台功能保留分析", "subtitle": "2026-08-12 猫娘"},
    {"layout": "section", "title": "第一章 背景"},
    {"layout": "content", "title": "第一梯队", "bullets": ["任务工单管理（6/7）", "- 统计、列表、查询全覆盖", "问题工单管理（5/7）"], "notes": "强调任务工单是核心刚需"},
    {"layout": "end", "title": "谢谢观看", "subtitle": "有问题随时喵我"},
]
r = T.tool_create_pptx(pptx_p, slides, confirmed=True)
ok &= check("生成成功", "已经生成" in r and os.path.exists(pptx_p), r[:60])

prs = Presentation(pptx_p)
ok &= check("页数=4", len(prs.slides) == 4)
# 逐页文字
texts = []
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            texts.append(sh.text_frame.text)
joined = "|".join(texts)
for kw in ("水务平台功能保留分析", "2026-08-12", "第一章", "第一梯队", "任务工单管理", "谢谢观看"):
    ok &= check(f"含「{kw}」", kw in joined)
ok &= check("二级要点(– 前缀)", "–" in joined and "统计、列表" in joined)
# 中文 eastAsia 字体
all_font_ok = True
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    rPr = run._r.rPr
                    ea = rPr.find(qn("a:ea")) if rPr is not None else None
                    if ea is None or ea.get("typeface") != "微软雅黑":
                        all_font_ok = False
ok &= check("全部 run 设 eastAsia 微软雅黑", all_font_ok)
# 演讲备注
notes_texts = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
ok &= check("notes 写进备注", any("核心刚需" in n for n in notes_texts))
# 读回自检工具
r2 = T._pptx_read(Path(pptx_p))
ok &= check("_pptx_read 读回", "第 1 页" in r2 and "第 4 页" in r2, r2[:40])
# 门禁/上限
r = T.tool_create_pptx(pptx_p, slides, confirmed=False)
ok &= check("未确认拒绝", "安全确认" in r)
r = T.tool_create_pptx(os.path.join(tmp, "x.pdf"), slides, confirmed=True)
ok &= check("非pptx拒绝", ".pptx" in r)
r = T.tool_create_pptx(os.path.join(tmp, "big.pptx"), slides * 31, confirmed=True)
ok &= check("超30页拒绝", "30 页" in r)
bad = [{"layout": "content", "title": "x", "bullets": [f"要点{i}" for i in range(9)]}]
r = T.tool_create_pptx(os.path.join(tmp, "b.pptx"), bad, confirmed=True)
ok &= check("要点超8拒绝", "8 条" in r)
shutil.rmtree(tmp, ignore_errors=True)
print("ALL PASS" if ok else "HAS FAILURES")
