# -*- coding: utf-8 -*-
"""验证 create_xlsx 工具（隔离 APPDATA，不碰真实用户数据）。"""
import os, sys, glob, shutil, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

tmp = tempfile.mkdtemp(prefix="catgirl_xlsx_")
os.environ["APPDATA"] = tmp

from backend.tools import tool_create_xlsx, run_tool, TOOL_IMPL

results = []
def check(name, cond, extra=""):
    results.append(cond)
    print(("✅" if cond else "❌"), name, ("  " + str(extra)) if extra else "")

check("工具已注册", "create_xlsx" in TOOL_IMPL)

# 1. confirmed 门禁
r = tool_create_xlsx(r"C:\out.xlsx", [{"name": "s", "rows": [["a"]]}], confirmed=False)
check("confirmed 门禁拦截", "安全确认" in r, r[:36])

# 2. 后缀校验
r = tool_create_xlsx(r"C:\out.txt", [{"name": "s", "rows": [["a"]]}], confirmed=True)
check("非 .xlsx 拒绝", "得是 .xlsx 结尾" in r, r[:30])

# 3. 表名重复
r = tool_create_xlsx(r"C:\out.xlsx", [{"name": "a", "rows": [["1"]]}, {"name": "a", "rows": [["2"]]}], confirmed=True)
check("表名重复拒绝", "表名重复" in r, r[:30])

# 4. 空 rows
r = tool_create_xlsx(r"C:\out.xlsx", [{"name": "a", "rows": []}], confirmed=True)
check("空 rows 拒绝", "没有数据行" in r, r[:30])

# 5. 正常创建（多 sheet + 样式）
out = os.path.join(tmp, "测试表.xlsx")
sheets = [
    {"name": "按星期分类", "header": ["周一", "周二"], "rows": [["张三", "李四"], ["王五", "赵六"]]},
    {"name": "学生明细", "header": ["姓名", "周一"], "rows": [["张三", "✓"], ["王五", ""]]},
]
r = tool_create_xlsx(out, sheets, confirmed=True)
check("多 sheet 创建成功", "已经生成" in r, r[:70])
check("话术含表名", "按星期分类" in r)

from openpyxl import load_workbook
from openpyxl.styles import Font
wb = load_workbook(out, data_only=True)
check("读回 sheet 数与顺序", wb.sheetnames == ["按星期分类", "学生明细"], str(wb.sheetnames))
ws = wb["按星期分类"]
check("表头正确", [ws.cell(1, c).value for c in (1, 2)] == ["周一", "周二"])
check("数据正确", [ws.cell(2, c).value for c in (1, 2)] == ["张三", "李四"])
check("数据正确2", [ws.cell(3, c).value for c in (1, 2)] == ["王五", "赵六"])
check("表头加粗样式", ws.cell(1, 1).font.bold is True)
check("列宽自适应", ws.column_dimensions["A"].width >= 4, ws.column_dimensions["A"].width)

# 6. 覆盖已有文件 → 备份到 backup 目录
sheets2 = [{"name": "新表", "rows": [["A"]]}]
r = tool_create_xlsx(out, sheets2, confirmed=True)
check("覆盖已有文件", "覆盖了原来的文件" in r, r[:70])
backups = glob.glob(os.path.join(tmp, "catgirl", "backup", "*.xlsx"))
check("旧版已备份", len(backups) >= 1, len(backups))

# 7. run_tool 分发 + 参数错误兜底
r = run_tool("create_xlsx", {"path": out, "sheets": [{"name": "分发", "rows": [["1"]]}], "confirmed": True})
check("run_tool 分发正常", "已经生成" in r, r[:40])
r = run_tool("create_xlsx", {"path": out, "sheets": "bad", "confirmed": True})
check("参数错误兜底不崩", "喵" in r and "traceback" not in r.lower(), r[:40])

# 8. 敏感文件拒绝
r = tool_create_xlsx(os.path.join(tmp, ".env.xlsx"), [{"name": "s", "rows": [["1"]]}], confirmed=True)
check("敏感文件拒绝", "敏感文件" in r, r[:30])

shutil.rmtree(tmp, ignore_errors=True)
ok = sum(1 for c in results if c)
print(f"\n通过 {ok} / {len(results)}")
sys.exit(0 if ok == len(results) else 1)
