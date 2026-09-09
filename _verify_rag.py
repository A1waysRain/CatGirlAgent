# 验证：backend/rag.py 知识库检索（RAG，2026-08-27 接入猫娘）
#  A) 切分：按 Markdown / Word 标题切块、Excel 表头分批 / 超长节按字符切（800/100 重叠）
#  B) 查询改写 rewrite_query：英文缩写展开（RAG→检索增强生成 等）
#  C) 索引：构建/持久化/幂等加载/文档变化自动重建/空知识库
#  D) 检索：命中带来源 / rewrite 命中 / 无关查询拦下(nomatch) / 空库 / 构建中
#  E) 工具分发：TOOL_IMPL 注册 / schema / run_tool 走通 / 各话术 / 参数错误
#  F) 懒加载：模块顶层不 import fastembed / warmup 不抛异常
#
# 全程 mock embedding（_fake_embed 用可控主题向量），不下载模型、不耗 token。
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from backend import rag  # noqa: E402

PASS = FAIL = 0
_TMPS = []
_REAL_SRC_ROOT = rag._SRC_ROOT
_REAL_DEFAULT_DOCS = list(rag.DEFAULT_DOCS)


def ok(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        print(f"  ❌ {name}  {extra}")


def make_tmp(prefix):
    t = tempfile.mkdtemp(prefix=prefix)
    _TMPS.append(t)
    return t


# ---- 可控伪 embedding：按主题词映射到独热向量（余弦 1=同主题 / 0=无关） ----
def _fake_embed(texts):
    vecs = []
    for t in texts:
        if "检索增强" in t or "RAG" in t or "向量库" in t:
            vecs.append([1.0, 0.0, 0.0, 0.0])
        elif "流式" in t or "SSE" in t or "推送" in t:
            vecs.append([0.0, 1.0, 0.0, 0.0])
        elif "聊天" in t or "会话" in t:
            vecs.append([0.0, 0.0, 1.0, 0.0])
        else:
            vecs.append([0.0, 0.0, 0.0, 1.0])
    return vecs


def reset_rag(tmp):
    """隔离重置 rag 模块状态：docs/index 指向临时目录、禁用默认语料复制、换伪向量。"""
    rag.DOCS_DIR = Path(tmp) / "docs"
    rag.INDEX_DIR = Path(tmp) / "index"
    rag.DEFAULT_DOCS = []
    rag._ready = False
    rag._building = False
    rag._index = {"vecs": None, "chunks": []}
    rag._embed = _fake_embed
    rag.DOCS_DIR.mkdir(parents=True, exist_ok=True)


def write_doc(tmp, name, text):
    p = Path(tmp) / "docs" / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")


def write_docx(tmp, name):
    from docx import Document
    p = Path(tmp) / "docs" / name
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("RAG Word 标题", level=1)
    doc.add_paragraph("检索增强生成把资料和问题一起交给模型。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "字段", "说明"
    table.cell(1, 0).text, table.cell(1, 1).text = "RAG", "检索增强生成"
    doc.save(str(p))


def write_xlsx(tmp, name, rows=205):
    from openpyxl import Workbook
    p = Path(tmp) / "docs" / name
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "销售数据"
    ws.append(["月份", "主题"])
    for i in range(rows):
        ws.append([f"2026-{i + 1:03d}", "检索增强生成"])
    wb.save(str(p))


print("\n0) 默认语料路径 / 空库首次同步")
missing = [rel for rel in _REAL_DEFAULT_DOCS if not (_REAL_SRC_ROOT / rel).is_file()]
ok("DEFAULT_DOCS 全部指向现有文件", not missing, str(missing))

tmp0 = make_tmp("rag_default_sync_")
reset_rag(tmp0)
fake_src = Path(tmp0) / "source"
default_rel = "资料/默认教材.md"
default_src = fake_src / default_rel
default_src.parent.mkdir(parents=True, exist_ok=True)
default_src.write_text("# 默认教材\n检索增强生成默认语料。", encoding="utf-8")
rag._SRC_ROOT = fake_src
rag.DEFAULT_DOCS = [default_rel]
rag._building = True
sync_state, sync_text = rag.search("检索增强生成是什么")
ok("直接检索也会先同步默认语料",
   sync_state is False and sync_text == "building" and (Path(tmp0) / "docs" / default_rel).is_file())
rag._building = False
ok("空 docs 首次 ensure 能同步并建索引", rag.ensure_index(block=True))
ok("默认语料已复制到 docs", (Path(tmp0) / "docs" / default_rel).is_file())
ok("同步后的默认语料已进入索引",
   any(c.get("source") == default_rel for c in rag._index["chunks"]), str(rag._index["chunks"]))
rag._SRC_ROOT = _REAL_SRC_ROOT
rag.DEFAULT_DOCS = list(_REAL_DEFAULT_DOCS)


print("\nA) 切分")
rag._init_dirs()
chunks = rag._split_document("## 标题A\n内容A第一行\n\n# 标题B\n内容B", "t.md")
ok("标题切块 A/B 两块", len(chunks) == 2)
ok("块 title 正确", chunks[0]["title"] == "## 标题A" and chunks[1]["title"] == "# 标题B")
ok("块文本带标题", chunks[0]["text"].startswith("## 标题A"))
long_text = "字" * 2000
segs = rag._by_chars(long_text)
ok("超长切 3 段", len(segs) == 3, f"got {len(segs)}")
ok("首段 800 字符", len(segs[0]) == 800)
overlap = segs[1][:100]
ok("相邻段 100 重叠", segs[0][-100:] == overlap)
ok("短文单段", rag._by_chars("短") == ["短"])

tmp_struct = make_tmp("rag_structured_")
reset_rag(tmp_struct)
write_docx(tmp_struct, "说明.docx")
write_xlsx(tmp_struct, "数据.xlsx")
ok("Word 标题和表格进入结构化 chunk", rag.ensure_index(block=True) and any(
    c["source"] == "说明.docx" and c["title"] == "RAG Word 标题" and "字段 | 说明" in c["text"]
    for c in rag._index["chunks"]), str(rag._index["chunks"]))
xlsx_chunks = [c for c in rag._index["chunks"] if c["source"] == "数据.xlsx"]
ok("Excel 保留工作表、表头并按行分批", len(xlsx_chunks) >= 3 and all(
    c["title"].startswith("表：销售数据") and "表头：月份 | 主题" in c["text"] for c in xlsx_chunks), str(xlsx_chunks[:1]))

print("\nB) 查询改写 rewrite_query")
ok("RAG→检索增强生成", rag.rewrite_query("RAG 是什么") == "检索增强生成 是什么")
ok("SSE→服务器推送事件", "服务器推送事件" in rag.rewrite_query("什么是SSE"))
ok("LLM→大语言模型", "大语言模型" in rag.rewrite_query("LLM 怎么训练"))
ok("无缩写原样", rag.rewrite_query("今天几点") == "今天几点")

print("\nC) 索引构建/加载/重建")
tmp1 = make_tmp("rag_c_")
reset_rag(tmp1)
write_doc(tmp1, "a.md", "# RAG\n检索增强生成就是先检索再生成。向量库存 embedding。\n" * 3)
write_doc(tmp1, "b.md", "# 流式\n流式就是 SSE 服务器推送事件。\n" * 3)
write_doc(tmp1, "c.md", "# 聊天\n猫娘聊天用 function calling。\n" * 3)
ok("ensure_index(block=True) 成功", rag.ensure_index(block=True))
ok("ready=True", rag._ready)
ok("chunks ≥3", len(rag._index["chunks"]) >= 3)
for f in ("index.json", "vectors.npy", "sig.json"):
    ok(f"索引文件 {f} 生成", (Path(tmp1) / "index" / f).exists())
n0 = len(rag._index["chunks"])
ok("再次 ensure 幂等不重建", rag.ensure_index(block=True) and len(rag._index["chunks"]) == n0)
idx_mtime = (Path(tmp1) / "index" / "index.json").stat().st_mtime_ns
write_doc(tmp1, "a.md", "# RAG\n检索增强生成就是先检索再生成。向量库存 embedding。\n新增一段内容改变向量。\n" * 3)
rag.ensure_index(block=True)
ok("文档变化触发重建", (Path(tmp1) / "index" / "index.json").stat().st_mtime_ns != idx_mtime)
tmpE = make_tmp("rag_empty_")
reset_rag(tmpE)
ok("空知识库 ensure 不 ready", not rag.ensure_index(block=True))

print("\nD) 检索 search")
tmp2 = make_tmp("rag_d_")
reset_rag(tmp2)
write_doc(tmp2, "a.md", "# RAG\n检索增强生成就是先检索再生成。\n" * 3)
write_doc(tmp2, "b.md", "# 流式\n流式就是 SSE 服务器推送事件。\n" * 3)
rag.ensure_index(block=True)
s, ctx = rag.search("什么是检索增强生成")
ok("命中返回 True", s is True)
ok("context 带来源 a.md", "a.md" in ctx and "检索增强" in ctx)
s2, ctx2 = rag.search("RAG 是什么")
ok("rewrite 后命中", s2 is True and "a.md" in ctx2)
s3, ctx3 = rag.search("今天天气怎么样")
ok("无关查询拦下 nomatch", s3 is True and ctx3 == "nomatch", f"got {ctx3[:40]}")
reset_rag(tmpE)
s4, ctx4 = rag.search("x")
ok("空库返回 empty", s4 is True and ctx4 == "empty")
reset_rag(tmp2)
rag._building = True
s5, ctx5 = rag.search("x")
ok("构建中返回 building", s5 is False and ctx5 == "building")
rag._building = False

print("\nE) 工具分发 rag_query")
from backend.tools import TOOL_IMPL, TOOL_SCHEMAS, run_tool  # noqa: E402
ok("TOOL_IMPL 注册 rag_query", "rag_query" in TOOL_IMPL)
sch = next((x for x in TOOL_SCHEMAS if x["function"]["name"] == "rag_query"), None)
ok("schema 存在且 query 必填", sch is not None and "query" in sch["function"]["parameters"]["required"])
reset_rag(tmp2)
rag.ensure_index(block=True)
out = run_tool("rag_query", {"query": "什么是检索增强生成", "top_k": 2})
ok("run_tool 走通返回片段", "a.md" in out and "检索增强" in out)
reset_rag(tmpE)
out_empty = run_tool("rag_query", {"query": "x"})
ok("空知识库引导话术", "还是空的" in out_empty)
reset_rag(tmp2)
rag.ensure_index(block=True)
out_nomatch = run_tool("rag_query", {"query": "今天天气怎么样"})
ok("没找到话术", "没找到" in out_nomatch)
def _boom(q, k):
    raise RuntimeError("boom")
rag.search = _boom
out_err = run_tool("rag_query", {"query": "x"})
ok("异常转话术", "检索失败" in out_err)
out_bad = run_tool("rag_query", {})
ok("缺参数 TypeErr 话术", "参数不对" in out_bad)

print("\nF) 懒加载")
ok("模块顶层未引 fastembed", "fastembed" not in sys.modules)
reset_rag(make_tmp("rag_f_"))
rag.warmup()
ok("warmup 不抛异常", True)

print(f"\n===== 结果：PASS {PASS} / FAIL {FAIL} =====")
for t in _TMPS:
    shutil.rmtree(t, ignore_errors=True)
sys.exit(1 if FAIL else 0)
