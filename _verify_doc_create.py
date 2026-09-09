# -*- coding: utf-8 -*-
"""隔离 APPDATA 验证 append_file 新建 docx/xlsx/pptx + 既有文件追加回归。"""
import os, sys, tempfile, shutil
tmp = tempfile.mkdtemp(prefix="catgirl_test_")
os.environ["APPDATA"] = tmp
sys.path.insert(0, ".")
import backend.tools as T
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

def check(label, cond, extra=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {label}" + (f"  {extra}" if extra else ""))
    return cond

ok = True
# ① 新建 docx（桌面裸文件名 → 但这里改绝对路径到临时目录）
docx_p = os.path.join(tmp, "功能保留分析.docx")
r = T.tool_append_file(docx_p, "第一行\n第二行内容", confirmed=True)
ok &= check("新建docx", "新建 Word 文档" in r and os.path.exists(docx_p), r[:40])
doc = Document(docx_p)
ok &= check("docx内容", [p.text for p in doc.paragraphs] == ["第一行", "第二行内容"])

# ② 新建 xlsx
xlsx_p = os.path.join(tmp, "清单.xlsx")
r = T.tool_append_file(xlsx_p, "甲\n乙", confirmed=True)
ok &= check("新建xlsx", "新建 Excel" in r and os.path.exists(xlsx_p), r[:40])
wb = load_workbook(xlsx_p)
ok &= check("xlsx内容", wb.active["A1"].value == "甲" and wb.active["A2"].value == "乙")

# ③ 新建 pptx
pptx_p = os.path.join(tmp, "汇报.pptx")
r = T.tool_append_file(pptx_p, "要点1\n要点2", confirmed=True)
ok &= check("新建pptx", "新建 PPT" in r and os.path.exists(pptx_p), r[:40])
prs = Presentation(pptx_p)
ok &= check("pptx有1页", len(prs.slides) == 1)

# ④ 既有 docx 追加（不新建、保留原内容）
r = T.tool_append_file(docx_p, "追加段", confirmed=True)
doc2 = Document(docx_p)
ok &= check("既有docx追加", len(doc2.paragraphs) == 3 and doc2.paragraphs[2].text == "追加段", r[:30])

# ⑤ 既有 txt 追加回归
txt_p = os.path.join(tmp, "测试.txt")
open(txt_p, "w", encoding="utf-8").write("旧内容\n")
r = T.tool_append_file(txt_p, "新内容", confirmed=True)
ok &= check("txt追加", open(txt_p, encoding="utf-8").read() == "旧内容\n新内容\n", r[:30])

# ⑥ confirmed 门禁
r = T.tool_append_file(os.path.join(tmp, "x.docx"), "内容", confirmed=False)
ok &= check("未确认拒绝", "安全确认" in r)

# ⑦ 非文档非文本后缀拒绝
r = T.tool_append_file(os.path.join(tmp, "bad.exe"), "x", confirmed=True)
ok &= check("坏扩展拒绝", "不能改" in r)

shutil.rmtree(tmp, ignore_errors=True)
print("ALL PASS" if ok else "HAS FAILURES")
